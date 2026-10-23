#!/usr/bin/env python3
"""Create v55 from v54 with evidence-aligned experiments and automatic numbering."""
from pathlib import Path
import re
import uuid

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

SRC = Path("/Users/mthang1201/Downloads/Context_Aware_Profile_Selection_Extension_JupyterHub_v54.pptx")
DST = Path("docs/slides/Context_Aware_Profile_Selection_Extension_JupyterHub_v55.pptx")
NAVY = RGBColor(7, 36, 75)
BLUE = RGBColor(37, 99, 235)
ORANGE = RGBColor(244, 125, 32)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)
PALE = RGBColor(246, 249, 254)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(185, 28, 28)
GREEN = RGBColor(22, 101, 52)


def delete_slide(prs, index):
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def clear_slide(slide):
    tree = slide.shapes._spTree
    for child in list(tree)[2:]:
        tree.remove(child)


def textbox(slide, x, y, w, h, text, size=14, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Arial", margin=0.06):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold if i == 0 else False
        p.font.color.rgb = color
        p.space_after = Pt(5)
        p.alignment = align
    return shape


def box(slide, x, y, w, h, fill=WHITE, line=RGBColor(203, 213, 225)):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(1)
    return s


def header(slide, title, subtitle, tag):
    clear_slide(slide)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PALE
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.15))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    side = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(.15), Inches(.10), Inches(7.35))
    side.fill.solid(); side.fill.fore_color.rgb = ORANGE; side.line.fill.background()
    textbox(slide, .55, .35, 11.9, .45, title, 22, NAVY, True)
    textbox(slide, .58, .84, 11.9, .35, subtitle, 11, MUTED)
    textbox(slide, 11.9, .38, .8, .25, tag, 9, BLUE, True, PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.58), Inches(7.04), Inches(12.1), Inches(.012))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(203, 213, 225); line.line.fill.background()


def section(slide, x, y, w, title, body, h=4.7, accent=BLUE, body_size=12):
    box(slide, x, y, w, h)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.07), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, x+.2, y+.16, w-.4, .4, title, 15, NAVY, True)
    textbox(slide, x+.2, y+.66, w-.4, h-.82, body, body_size, INK)


def make_e12(slide):
    header(slide, "E1 + E2 | Intent input to recommended profile", "Offline algorithm evaluation. Results are development evidence and do not support confirmatory claims.", "E1 / E2")
    section(slide, .6, 1.38, 3.8, "INPUT · workload intent", "83 bilingual prompt variants\n18 semantic workload families\n27 English + 56 Vietnamese\n\nWorkload intent covers task, libraries, data size, CPU/RAM needs, GPU constraints, and wording variants.\n\nLabels: acceptable profiles, images, and hard constraints.", 4.95, BLUE, 12)
    section(slide, 4.75, 1.38, 3.8, "METHOD · P1 vs P2", "Same frozen cases and candidate catalog for both systems.\n\nP1: rule-based baseline.\nP2: local structured-intent extraction, sparse+dense retrieval, reciprocal-rank fusion, then feasibility constraints.\n\nE1 scores recommendation quality. E2 compares recommendation stability across equivalent prompt variants.", 4.95, ORANGE, 12)
    section(slide, 8.9, 1.38, 3.8, "OUTPUT · observed development results", "E1 joint accept: P1 32/83 (38.6%); P2 39/83 (47.0%). P2 fallback: 28/83 (33.7%). Counts recomputed from paired raw records.\n\nE2 macro family SRR: P1 56.8% [31.8–79.5]; P2 72.0% [55.3–87.9]. Some variant strata are N/A.\n\nDevelopment evidence only. Family is the unit; do not treat 83 variants as independent families or claim confirmatory superiority.", 4.95, GREEN, 10)
    textbox(slide, .72, 6.48, 11.8, .35, "Environment: Apple M2 / 16 GB host as reported by researcher; CPU-only local evaluation, no GPU. E1/E2 do not require OrbStack or Kubernetes.", 10, MUTED)


def make_e4(slide):
    header(slide, "E4 | Resource fit under real Kubernetes Pod execution", "Observed OrbStack run: 800 Pod trials across four provisioning policies. One ARM64 node; results are specific to this testbed.", "E4")
    section(slide, .6, 1.38, 3.8, "INPUT · workloads + policies", "16 workload instances repeated across the frozen plan.\n\nConditions: P1 Catalog, P2 Catalog, P2 Dynamic, Static Large.\n\nStress workloads include streaming statistics, time-series, joins, ML / matrix computation, and related bounded tasks.\n\nEach condition: 200 trials.", 4.95, BLUE, 11.5)
    section(slide, 4.75, 1.38, 3.8, "ENVIRONMENT · measured", "OrbStack Kubernetes v1.33.9+orb1\nARM64 single node; 8 CPU allocatable and about 7.8 GiB node memory.\nDocker 29.4.0; Linux kernel 7.0.5-orbstack; cgroups v2.\n\nHost: MacBook M2, 16 GB unified memory (researcher-reported). No physical GPU; E4 trials requested CPU only.", 4.95, ORANGE, 11)
    section(slide, 8.9, 1.38, 3.8, "OUTPUT · observed outcomes", "Success / OOM: P1 90% / 20; P2 Catalog 100% / 0; P2 Dynamic 95% / 10; Static Large 100% / 0.\n\nMean peak memory: 16.5, 174.8, 88.5, 174.8 MiB respectively.\n\nP1’s low memory reflects early OOM exits. Dynamic reduced mean memory, with 10 OOMs. No multi-tenant cluster-capacity result was executed.", 4.95, GREEN, 10.5)
    textbox(slide, .72, 6.48, 11.8, .35, "Output files: raw Pod lifecycle and cgroup measurements, trial plan, environment provenance, and condition summaries. Treat this as one-cluster evidence.", 10, MUTED)


def make_e5(slide):
    header(slide, "E5 | Image capability checks and OCI layer storage", "Two separate measurements: selected-image functional probes and registry manifest layer accounting.", "E5")
    section(slide, .6, 1.38, 3.8, "INPUT · approved catalog", "Four digest-pinned Jupyter images: minimal, SciPy, PyTorch, and TensorFlow.\n\nFunctional run consumes a frozen offline recommendation set (83 cases per system) and checks declared capabilities against bounded in-container probes.\n\nStorage run inspects the four linux/arm64 OCI manifests.", 4.95, BLUE, 11.5)
    section(slide, 4.75, 1.38, 3.8, "ENVIRONMENT · Docker on M2", "Functional probes ran with Docker on macOS ARM64 via OrbStack-compatible Docker runtime.\n\nStorage evidence was collected from real registry manifests for linux/arm64; it is not a local disk-usage benchmark.\n\nHost reported: Apple M2, 16 GB unified memory. GPU execution was not tested.", 4.95, ORANGE, 11)
    section(slide, 8.9, 1.38, 3.8, "OUTPUT · bounded findings", "Functional probes passed: P1 65/65 eligible; P2 76/76 eligible. Catalog capability coverage: 58/83 (69.9%) P1; 70/83 (84.3%) P2.\n\nFour-image storage: 8.247 GB unique layers vs 11.393 GB logical sum; 27.6% layer deduplication.\n\nCatalog scales 8 and 16: NOT EXECUTED (only four approved images). Development evidence only.", 4.95, GREEN, 10.5)
    textbox(slide, .72, 6.48, 11.8, .35, "E5 validates bounded image behavior and four-image layer reuse. It does not establish broad catalog scalability or GPU compatibility.", 10, MUTED)


def revise_simulation(slide, is_results=False):
    title = "Capacity simulation | Current evidence boundary" if is_results else "Capacity simulation | What the proposed model would test"
    subtitle = "Repository audit: large-cluster outcomes in the previous slides were not produced by a validated Protocol-v5 simulation run."
    header(slide, title, subtitle, "SIMULATION")
    if not is_results:
        section(slide, .8, 1.55, 5.7, "PROPOSED INPUTS", "Generate tenant intents from a declared workload mix, then pass each intent through the same P1/P2 recommendation path.\n\nFeed resulting CPU/RAM/GPU requests and image identities to a discrete-event admission model. Declare arrival rates, capacity, scheduling rules, and random seeds.", 4.9, BLUE, 12)
        section(slide, 6.85, 1.55, 5.7, "PROPOSED OUTPUTS", "Compare admitted / pending sessions, reserved-but-idle resources, OOM assumptions, and image pulls under stated scenarios.\n\nValidate model inputs against E1/E2 recommendations and E4 measurements. E5 image storage can inform image-pull/storage cost, but it is not a scheduler experiment.", 4.9, ORANGE, 12)
        textbox(slide, .9, 6.62, 11.5, .28, "Simulation is a separate capacity study. It should exercise E1/E2 recommendation inputs and E4 scheduling behavior; E5 contributes only image-related parameters where relevant.", 10, MUTED)
    else:
        section(slide, .8, 1.55, 5.7, "WHAT EXISTS IN THE REPOSITORY", "The tracked E4 capacity artifact reports status NOT_EXECUTED and reason NODE_CAPACITY_NOT_FROZEN.\n\nThe legacy script scripts/simulate_large_cluster.py is not a validated source for the 500 / 1,000-session claims shown in v54.\n\nNo reproducible simulation output supports the prior 58.4%, 240 pending, 815–920 served, or 3.1×–3.5× figures.", 4.9, RED, 11)
        section(slide, 6.85, 1.55, 5.7, "WHAT MUST HAPPEN BEFORE CLAIMING RESULTS", "Freeze capacity, arrivals, workload mix, policy behavior, seeds, and validation checks. Run the same P1/P2 recommender against declared workloads, then retain raw event traces and a versioned result package.\n\nUntil then, present this as planned future work. Keep the 800 observed E4 Pod trials separate from the simulation.", 4.9, GREEN, 11.5)
        textbox(slide, .9, 6.62, 11.5, .28, "The current E4 run is real Pod execution on one OrbStack node. It does not equal a multi-tenant cluster-capacity simulation.", 10, MUTED)


def add_auto_number(slide, index, total):
    # Remove old manually typed footer counters.
    for shp in list(slide.shapes):
        if getattr(shp, "has_text_frame", False) and re.fullmatch(r"\s*\d+\s*/\s*\d+\s*", shp.text or ""):
            el = shp._element
            el.getparent().remove(el)
    shape = slide.shapes.add_textbox(Inches(11.95), Inches(7.08), Inches(.75), Inches(.22))
    tf = shape.text_frame; tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    p.font.name = "Arial"; p.font.size = Pt(8); p.font.color.rgb = MUTED
    fld = OxmlElement("a:fld")
    fld.set("id", "{" + str(uuid.uuid4()) + "}")
    fld.set("type", "slidenum")
    rpr = OxmlElement("a:rPr"); rpr.set("lang", "en-US"); rpr.set("sz", "800")
    solid = OxmlElement("a:solidFill"); clr = OxmlElement("a:srgbClr"); clr.set("val", "475569"); solid.append(clr); rpr.append(solid)
    latin = OxmlElement("a:latin"); latin.set("typeface", "Arial"); rpr.append(latin)
    fld.append(rpr)
    tx = OxmlElement("a:t"); tx.text = f"{index:02d}"
    fld.append(tx); p._p.append(fld)
    run = p.add_run(); run.text = f" / {total:02d}"; run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = MUTED


def main():
    prs = Presentation(SRC)
    # Attached v54 slide positions 19-24: E1, E2, E4, E5, simulation setup, simulation results.
    make_e12(prs.slides[18])
    delete_slide(prs, 19)  # E1/E2 now share one detailed workflow slide.
    make_e4(prs.slides[19])
    make_e5(prs.slides[20])
    revise_simulation(prs.slides[21], False)
    revise_simulation(prs.slides[22], True)
    # Remove unsupported quantitative thesis conclusion while preserving conclusion slide structure.
    for shp in prs.slides[23].shapes:
        if getattr(shp, "has_text_frame", False) and shp.text.strip() == "Summary of the implemented bounded architecture, empirical findings (E1–E5), and large-scale cluster simulation.":
            shp.text = "Summary of the architecture and the evidence currently available for E1/E2, E4, and E5."
        if getattr(shp, "has_text_frame", False) and shp.text.strip().startswith("DETERMINISTIC LOCAL NLP"):
            shp.text = "LOCAL NLP EXTRACTOR  ·  800 OBSERVED ORBSTACK POD TRIALS  ·  E5 FUNCTIONAL + STORAGE MEASUREMENTS"
        if getattr(shp, "has_text_frame", False) and any(x in shp.text for x in ("3.1×", "5.56x", "1,000-session", "58.4%")):
            shp.text = "E1/E2: Development benchmark of 83 prompt variants across 18 workload families. E4: 800 observed Pod trials on one OrbStack node. E5: 141 eligible functional probes passed and four-image OCI layer reuse measured. Multi-tenant capacity simulation remains future work."
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        add_auto_number(slide, i, total)
    prs.save(DST)
    print(f"Wrote {DST} ({total} slides)")


if __name__ == "__main__":
    main()
