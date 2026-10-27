#!/usr/bin/env python3
"""Sync presentation slides and presenter notes to resolve self-contradictions in v62."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

REPO_ROOT = Path("/Users/mthang1201/Documents/datn/intent-spawner")
PPTX_PATH = REPO_ROOT / "docs/slides/Context_Aware_Profile_Selection_Extension_JupyterHub_v62.pptx"

def main():
    print(f"Loading {PPTX_PATH}...")
    prs = Presentation(str(PPTX_PATH))

    # =========================================================================
    # 1. Slide 9 (Index 8): Table & Notes sync
    # =========================================================================
    s9 = prs.slides[8]
    for sp in s9.shapes:
        if sp.has_table:
            for row in sp.table.rows:
                crit = row.cells[0].text.strip()
                if "1. Hallucination Risk" in crit:
                    # Update cell 3
                    row.cells[2].text = (
                        "E5 evaluated 16 allowlisted, digest-pinned images across scales 4, 8, and 16. "
                        "All candidate pairs strictly pass PolicyValidator."
                    )
    if s9.has_notes_slide and s9.notes_slide.notes_text_frame:
        s9.notes_slide.notes_text_frame.text = (
            "Danh mục được phê duyệt và bộ lọc cứng giới hạn các lựa chọn hợp lệ, ngăn chặn hoàn toàn rủi ro hallucination. "
            "Trên 471 ca kiểm thử, P2 LLM đạt Candidate Accuracy@1 cao nhất (27,39%), P2 Local dẫn đầu vượt trội ở Candidate Accuracy@3 (73,46% so với 26,33% của P1), "
            "trong khi P1 dẫn đầu JointAccept@1 (84,29% so với 79,19% của P2 LLM và 66,03% của P2 Local). "
            "Độ trễ trung bình của P2 Local là 2,28 ms (sub-5ms) chạy offline 0$ chi phí, trong khi P2 LLM mất trung bình 1.508 ms qua API ngoài."
        )

    # =========================================================================
    # 2. Slide 18 (Index 17): Evaluation design - Candidate catalog text & Notes
    # =========================================================================
    s18 = prs.slides[17]
    for sp in s18.shapes:
        if sp.has_text_frame:
            txt = sp.text_frame.text
            if "Candidate catalog" in txt and "four-image" in txt:
                sp.text_frame.text = (
                    "Candidate catalog  |Offline test: Small/Medium/Large profiles; E4/E5 evaluate an administrator-approved 16-image operational catalog."
                )
    if s18.has_notes_slide and s18.notes_slide.notes_text_frame:
        s18.notes_slide.notes_text_frame.text = (
            "Thiết kế đánh giá gồm các nhánh bằng chứng thực nghiệm rõ ràng: E1 và E2 sử dụng tập holdout 471 ca (157 họ tách biệt) từ DS-1000 và ARCADE; "
            "Kaggle gồm 1.000 truy vấn thực tế cho đánh giá truy hồi hybrid (Hit@1 90%, MRR 0,9167). "
            "E4 gồm 800 lượt chạy Pod trực tiếp trên Kubernetes OrbStack (0 OOM với P2 Catalog) và bài toán đóng gói xác định (bin-packing) trên năng lực allocatable đã xác thực giúp giảm 50% số node. "
            "Mô phỏng 50 node 30 seed bổ sung góc nhìn vĩ mô về triệt tiêu hàng đợi và chống lãng phí GPU. "
            "E5 đánh giá năng lực probe và lưu trữ OCI trên toàn bộ danh mục 16 image mở rộng."
        )

    # =========================================================================
    # 3. Slide 20 (Index 19): E1 workflow - Remove conflicting note
    # =========================================================================
    s20 = prs.slides[19]
    for sp in s20.shapes:
        if sp.has_text_frame:
            txt = sp.text_frame.text.strip()
            if txt.startswith("Latency  |not reported for this test"):
                sp.text_frame.text = (
                    "Latency  |Measured end-to-end: P2 Local 2.28 ms mean (offline, sub-5ms) vs P2 LLM 1,508 ms mean (external API)"
                )
    if s20.has_notes_slide and s20.notes_slide.notes_text_frame:
        s20.notes_slide.notes_text_frame.text = (
            "Quy trình E1 đánh giá toàn diện cả ba cấu hình: P1 Baseline (luật từ khóa), P2 Local NLP (trích xuất xác định + truy hồi kết hợp BM25/Dense/RRF), "
            "và P2 LLM adapter (Gemini 3.5 Flash-Lite) trên tập kiểm thử 471 ca thuộc 157 nhóm workload độc lập. "
            "Mỗi bản ghi kiểm thử lưu trữ đầy đủ ngữ cảnh mã nguồn, kích thước dữ liệu và nhãn tài nguyên đo thực tế bằng Linux cgroups v2."
        )

    # =========================================================================
    # 4. Slide 21 (Index 20): E1 held-out results - Add P2 LLM to chart & Notes
    # =========================================================================
    s21 = prs.slides[20]
    
    # Legend: Add P2 LLM box & text
    # P1: rect at 4343400, text at 4517136 (color #2F75B5)
    # P2 Local: rect at 5486400, text at 5660136 (color #2A9D8F)
    # Add P2 LLM: rect at 6629400, text at 6803136 (color #E76F51)
    rect_llm = s21.shapes.add_shape(MSO_SHAPE.RECTANGLE, 6629400, 1170432, 128016, 128016)
    rect_llm.fill.solid()
    rect_llm.fill.fore_color.rgb = RGBColor(0xE7, 0x6F, 0x51)
    rect_llm.line.fill.background()

    tb_llm_legend = s21.shapes.add_textbox(6803136, 1124712, 1051560, 219456)
    tf_llm = tb_llm_legend.text_frame
    tf_llm.word_wrap = False
    p_llm = tf_llm.paragraphs[0]
    p_llm.text = "P2 LLM"
    p_llm.font.size = Pt(8.5)
    p_llm.font.name = "Sarabun"
    p_llm.font.color.rgb = RGBColor(0x06, 0x2B, 0x4F)

    # Chart scale mapping:
    # 0% -> 2423160, 100% -> 6080760. Delta = 3657600 (36576 per 1%)
    def pct_to_x(pct):
        return int(2423160 + (pct / 100.0) * 3657600)

    # Add 4 orange dots and text boxes for P2 LLM:
    # 1. JointAccept@1: 79.2% -> y=2130552 (text y=2084832)
    # 2. Candidate Acc@1: 27.4% -> y=2971800 (text y=2926080)
    # 3. Candidate Acc@3: 31.4% -> y=3813048 (text y=3767328)
    # 4. Under-prov: 6.4% -> y=4654296 (text y=4608576)
    llm_metrics = [
        (79.193, 2130552, 2084832, "79.2%"),
        (27.389, 2971800, 2926080, "27.4%"),
        (31.423, 3813048, 3767328, "31.4%"),
        (6.369, 4654296, 4608576, "6.4%"),
    ]

    for pct, dot_y, text_y, label_str in llm_metrics:
        dot_x = pct_to_x(pct)
        dot = s21.shapes.add_shape(MSO_SHAPE.OVAL, dot_x, dot_y, 91440, 91440)
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor(0xE7, 0x6F, 0x51)
        dot.line.fill.background()

        tb_val = s21.shapes.add_textbox(6199632, text_y, 1417320, 182880)
        tf_val = tb_val.text_frame
        tf_val.word_wrap = False
        p_val = tf_val.paragraphs[0]
        p_val.text = label_str
        p_val.font.size = Pt(8.5)
        p_val.font.name = "Sarabun"
        p_val.font.color.rgb = RGBColor(0x06, 0x2B, 0x4F)

    # Update Slide 21 Notes
    if s21.has_notes_slide and s21.notes_slide.notes_text_frame:
        s21.notes_slide.notes_text_frame.text = (
            "Trên 471 ca test của Benchmark v1, cả ba cấu hình thể hiện rõ sự đánh đổi: "
            "P1 Baseline đạt JointAccept@1 cao nhất (84,29%) nhờ xu hướng cấp dư tài nguyên bảo thủ (over-provision 62,84%). "
            "P2 Local NLP đạt Candidate Accuracy@3 vượt trội (73,46% so với 26,33% của P1), hoàn thành gợi ý trong 2,28 ms với chi phí 0$. "
            "P2 LLM adapter (Gemini 3.5 Flash-Lite) dẫn đầu ở độ chính xác khắt khe Candidate Accuracy@1 (27,39%) và Profile Accuracy (35,03%), "
            "đạt JointAccept@1 79,19% và giảm tỷ lệ cấp thừa xuống thấp nhất (58,60%), với độ trễ gọi API ngoài 1.508 ms."
        )

    # =========================================================================
    # 5. Slide 24 & 25 (Index 23, 24): E4 live Pods
    # =========================================================================
    s24 = prs.slides[23]
    if s24.has_notes_slide and s24.notes_slide.notes_text_frame:
        s24.notes_slide.notes_text_frame.text = (
            "Quy trình thực nghiệm E4 đánh giá 800 lượt chạy Pod trực tiếp trên Kubernetes OrbStack sử dụng 16 workload đóng băng "
            "(chia đều theo các họ streaming statistics, data joins, ML/matrix compute, bounded tasks). "
            "Bốn chính sách cấp phát được so sánh gồm P1 Catalog, P2 Catalog, P2 Dynamic và Static Large (mỗi chính sách 200 lượt)."
        )

    # =========================================================================
    # 6. Slide 26 & 27 (Index 25, 26): E5 image & storage
    # =========================================================================
    s26 = prs.slides[25]
    for sp in s26.shapes:
        if sp.has_text_frame:
            txt = sp.text_frame.text
            if "Catalog\nfour images" in txt:
                sp.text_frame.text = "Catalog\n16 approved digest-pinned images across scales 4, 8, 16"
            elif "Catalog coverage -- P1 69.9% (58/83)" in txt:
                sp.text_frame.text = (
                    "Catalog coverage (dev split) -- P1 69.9% (58/83) · P2 84.3% (70/83)\n"
                    "Functional probes -- P1 65/65 passed · P2 76/76 passed (100%)\n"
                    "Storage -- Evaluated across scales 4, 8, and 16 images via real Docker registry manifest inspection"
                )
    if s26.has_notes_slide and s26.notes_slide.notes_text_frame:
        s26.notes_slide.notes_text_frame.text = (
            "Quy trình E5 đo đạc năng lực ảnh container và mức tiết kiệm lưu trữ OCI: "
            "Độ phủ danh mục và kiểm thử probe chức năng được thực hiện trên tập development 83 ca đóng băng (đạt 100% probe thành công: 65/65 với P1 và 76/76 với P2). "
            "Thực nghiệm lưu trữ OCI được thực thi trên toàn bộ 16 image đã ghim digest của hệ sinh thái Jupyter Docker Stacks qua API manifest của registry Docker/Quay."
        )

    s27 = prs.slides[26]
    for sp in s27.shapes:
        if sp.has_text_frame:
            txt = sp.text_frame.text
            if "Coverage: P1 69.9% (58/83)" in txt:
                sp.text_frame.text = (
                    "Development split coverage (83 cases): P1 69.9% (58/83) · P2 84.3% (70/83) · probes passed 65/65 and 76/76 (100%)."
                )
    if s27.has_notes_slide and s27.notes_slide.notes_text_frame:
        s27.notes_slide.notes_text_frame.text = (
            "Kết quả E5 xác nhận năng lực ảnh và hiệu quả chia sẻ layer OCI: "
            "Trên tập development 83 ca, P2 đạt độ phủ 84,3% (so với 69,9% của P1) và 100% probe chức năng vượt qua kiểm thử. "
            "Về lưu trữ OCI thực tế, khi quy mô catalog mở rộng từ 4 lên 8 và 16 image, tỷ lệ tiết kiệm layer tăng đơn điệu từ 27,61% (2,93 GiB) "
            "lên 30,55% (5,34 GiB) và đạt 37,06% (tiết kiệm 12,48 GiB trên tổng 33,69 GiB dung lượng logic), khẳng định tính khả thi của catalog 16 image."
        )

    # =========================================================================
    # 7. Slide 28 & 29 (Index 27, 28): E4 Capacity & Simulation
    # =========================================================================
    s28 = prs.slides[27]
    for sp in s28.shapes:
        if sp.has_text_frame:
            txt = sp.text_frame.text
            if "Balanced mix of 20 distinct empirical E4 workload families" in txt:
                sp.text_frame.text = (
                    "Balanced mix of 16 frozen empirical E4 workload instances across compute, joins, ML, and bounded streaming tasks."
                )
    if s28.has_notes_slide and s28.notes_slide.notes_text_frame:
        s28.notes_slide.notes_text_frame.text = (
            "Đánh giá sức chứa chính thức E4 giải quyết bài toán đóng gói đa chiều (multi-dimensional bin-packing) "
            "trên tài nguyên allocatable đã được xác thực từ Kubernetes node thực tế (8 vCPU, 7,81 GiB RAM). "
            "Sử dụng 16 workload đóng băng từ thực nghiệm E4, P2 Catalog đóng gói toàn bộ 20 phiên vào 2 node mà không xảy ra rủi ro OOM, "
            "giúp giảm 50% số lượng node yêu cầu so với Static Large (cần 4 node) và tăng gấp đôi mật độ phục vụ (+100%)."
        )

    s29 = prs.slides[28]
    if s29.has_notes_slide and s29.notes_slide.notes_text_frame:
        s29.notes_slide.notes_text_frame.text = (
            "Mô phỏng cụm quy mô lớn (50 node, 1.000 vCPU, 2.000 GB RAM, 100 GPU) qua 30 seed và 24.000 phiên "
            "sử dụng dữ liệu telemetry thực tế từ cgroups v2 của E4. Kết quả cho thấy P2 duy trì tỷ lệ OOM 0,00%, "
            "triệt tiêu hoàn toàn hàng đợi chờ (0 Pods chờ so với đỉnh 63,6 Pods ở B0), và xóa bỏ hiện tượng nghẽn GPU bị giam giữ (GPU stranding giảm từ 69,47% về 0,00%)."
        )

    # Save to both v62 and output
    prs.save(str(PPTX_PATH))
    print(f"Successfully updated and saved {PPTX_PATH}")

    # Also update presentation_output pptx
    out_path = REPO_ROOT / "presentation_output/Context_Aware_Profile_Selection_Extension_JupyterHub_v49.pptx"
    if out_path.exists():
        prs.save(str(out_path))
        print(f"Also synced to {out_path}")

if __name__ == "__main__":
    main()
